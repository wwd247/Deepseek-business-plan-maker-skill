#!/usr/bin/env python3
"""build_docx.py — Convert structured business-plan JSON into a .docx."""
import json, sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    sys.stderr.write("Error: pip install python-docx\n")
    sys.exit(2)


def create_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for level, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[level]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    return doc


def add_footer(doc, company_name):
    footer = doc.sections[0].footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if company_name:
        run = para.add_run(f"{company_name}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_table(doc, spec):
    headers = spec.get("headers", [])
    rows = spec.get("rows", [])
    if not headers and not rows:
        return
    ncols = len(headers) if headers else (len(rows[0]) if rows else 0)
    if ncols == 0:
        return
    table = doc.add_table(rows=0, cols=ncols)
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        table.style = "Table Grid"
    if headers:
        hdr_cells = table.add_row().cells
        for i, text in enumerate(headers):
            hdr_cells[i].text = ""
            run = hdr_cells[i].paragraphs[0].add_run(str(text))
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            if i < ncols:
                cells[i].text = str(text)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        try:
            doc.add_paragraph(str(item), style="List Bullet")
        except KeyError:
            doc.add_paragraph(f"* {item}")


def add_section(doc, section):
    level = section.get("level", 1)
    if level < 1 or level > 3:
        level = 1
    if section.get("page_break_before"):
        doc.add_page_break()
    if section.get("heading"):
        doc.add_heading(str(section["heading"]), level=level)
    for para in section.get("body", []):
        if para:
            doc.add_paragraph(str(para))
    if section.get("bullets"):
        add_bullets(doc, section["bullets"])
    if section.get("table"):
        add_table(doc, section["table"])
    for nested in section.get("subsections", []):
        add_section(doc, nested)


def build(doc, data):
    title = data.get("title", "Business Plan")
    subtitle = data.get("subtitle")
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tp.add_run(str(title))
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    if subtitle:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sp.add_run(str(subtitle))
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()
    for section in data.get("sections", []):
        add_section(doc, section)


def main(argv):
    if len(argv) != 3:
        sys.stderr.write("Usage: build_docx.py <input.json | -> <output.docx>\n")
        return 1
    input_arg, output_path = argv[1], argv[2]
    try:
        raw = sys.stdin.read() if input_arg == "-" else Path(input_arg).read_text(encoding="utf-8")
        data = json.loads(raw)
    except json.JSONDecodeError as e:
        sys.stderr.write(f"Error: invalid JSON — {e}\n")
        return 1
    except OSError as e:
        sys.stderr.write(f"Error: could not read input — {e}\n")
        return 1
    if not isinstance(data, dict) or "sections" not in data:
        sys.stderr.write("Error: input must be an object with 'sections'.\n")
        return 1
    doc = create_document()
    add_footer(doc, data.get("company"))
    build(doc, data)
    try:
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
    except OSError as e:
        sys.stderr.write(f"Error: could not write output — {e}\n")
        return 1
    print(f"Wrote {output_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
